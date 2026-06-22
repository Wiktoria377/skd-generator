"""
Serwis parsowania szablonu pozwu SKD.
Analizuje dokument .docx i identyfikuje wszystkie pola do wypełnienia.
"""
import re
from dataclasses import dataclass
from docx import Document
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from typing import Optional


@dataclass
class PlaceholderLocation:
    paragraph_index: int
    run_indices: list
    placeholder_text: str
    field_id: str
    description: str
    category: str


PLACEHOLDER_PATTERNS = [
    (r'_{3,}', 'underscore'),
    (r'…{2,}', 'dots'),
    (r'\.{4,}', 'dots_ascii'),
]

FIELD_CONTEXT_MAP = {
    # (context_before, context_after) -> (field_id, description, category)
    ('dnia', 'r.'): ('data_pozwu', 'Data złożenia pozwu', 'dane_procesowe'),
    ('Sąd', ''): ('sad', 'Nazwa sądu', 'dane_procesowe'),
    ('Powód:', ''): ('powod_imie_nazwisko', 'Imię i nazwisko powoda', 'dane_stron'),
    ('Pozwany:', ''): ('pozwany_nazwa', 'Nazwa pozwanego (bank)', 'dane_stron'),
    ('WPS:', ''): ('wps', 'Wartość przedmiotu sporu', 'dane_procesowe'),
    ('Opłata:', ''): ('oplata', 'Opłata sądowa', 'dane_procesowe'),
    ('od pozwanego', 'z siedzibą'): ('pozwany_nazwa_petitum', 'Nazwa pozwanego banku w petitum', 'dane_stron'),
    ('siedzibą w', 'na rzecz'): ('pozwany_siedziba', 'Siedziba pozwanego', 'dane_stron'),
    ('na rzecz powoda', 'kwoty'): ('powod_imie_petitum', 'Imię powoda w petitum', 'dane_stron'),
    ('kwoty', 'złotych'): ('kwota_roszczenia', 'Kwota roszczenia', 'dane_finansowe'),
    ('słownie:', ')'): ('kwota_roszczenia_slownie', 'Kwota roszczenia słownie', 'dane_finansowe'),
    ('od dnia', 'do dnia'): ('data_wymagalnosci', 'Data wymagalności', 'dane_procesowe'),
    ('Umowy', 'z dnia'): ('nazwa_umowy', 'Nazwa/typ umowy', 'dane_umowy'),
    ('z dnia', 'zawartej'): ('data_zawarcia_umowy', 'Data zawarcia umowy', 'dane_umowy'),
    ('pomiędzy', 'a'): ('powod_pelne_dane', 'Pełne dane powoda', 'dane_stron'),
    ('a', ', zastrzegając'): ('pozwany_pelne_dane', 'Pełne dane pozwanego', 'dane_stron'),
    ('jedynie na dzień', 'r.'): ('data_naliczenia', 'Data naliczenia roszczenia', 'dane_procesowe'),
    ('po dniu', 'r.,'): ('data_po_ktorej', 'Data po której dalsze roszczenia', 'dane_procesowe'),
    ('z pozwanym (', ') w dniu'): ('nazwa_umowy_ustalenie', 'Nazwa umowy w ustaleniu', 'dane_umowy'),
    ('pożyczki w wysokości', 'zł'): ('kwota_pozyczki', 'Kwota pożyczki', 'dane_finansowe'),
    ('słownie:', ').'): ('kwota_pozyczki_slownie', 'Kwota pożyczki słownie', 'dane_finansowe'),
    ('spłacić', 'ratach'): ('liczba_rat', 'Liczba rat', 'dane_finansowe'),
    ('do', 'dnia każdego'): ('dzien_platnosci', 'Dzień płatności raty', 'dane_finansowe'),
    ('według', 'stopy procentowej'): ('typ_oprocentowania', 'Typ stopy procentowej', 'dane_finansowe'),
    ('wynosiła', '.'): ('oprocentowanie', 'Oprocentowanie', 'dane_finansowe'),
    ('pożyczka w kwocie', 'zł obejmować'): ('kwota_pozyczki_z_kosztami', 'Kwota pożyczki z kosztami', 'dane_finansowe'),
    ('tj.', '.'): ('opis_kosztow', 'Opis kosztów kredytowanych', 'dane_finansowe'),
    ('prowizja w kwocie', ', składki'): ('prowizja', 'Kwota prowizji', 'dane_finansowe'),
    ('ubezpieczeniowe w kwocie', ', zostały'): ('ubezpieczenie', 'Kwota ubezpieczenia', 'dane_finansowe'),
}


def parse_template(filepath: str) -> list[PlaceholderLocation]:
    """Parsuje szablon .docx i zwraca listę miejsc do wypełnienia."""
    doc = Document(filepath)
    placeholders = []
    field_counter = 0

    for p_idx, paragraph in enumerate(doc.paragraphs):
        full_text = paragraph.text
        if not full_text.strip():
            continue

        for pattern, ptype in PLACEHOLDER_PATTERNS:
            for match in re.finditer(pattern, full_text):
                start, end = match.start(), match.end()
                field_id, desc, cat = _identify_field(full_text, start, end, field_counter)
                run_indices = _find_run_indices(paragraph, start, end)

                placeholders.append(PlaceholderLocation(
                    paragraph_index=p_idx,
                    run_indices=run_indices,
                    placeholder_text=match.group(),
                    field_id=field_id,
                    description=desc,
                    category=cat,
                ))
                field_counter += 1

    return placeholders


def _identify_field(text: str, start: int, end: int, counter: int) -> tuple[str, str, str]:
    """Identyfikuje pole na podstawie kontekstu tekstowego."""
    context_before = text[max(0, start - 40):start].strip()
    context_after = text[end:end + 40].strip()

    for (ctx_b, ctx_a), (fid, desc, cat) in FIELD_CONTEXT_MAP.items():
        if ctx_b and ctx_b.lower() in context_before.lower():
            if not ctx_a or ctx_a.lower() in context_after.lower():
                return fid, desc, cat

    return f'pole_{counter}', f'Pole do uzupełnienia (kontekst: ...{context_before[-20:]} [___] {context_after[:20]}...)', 'inne'


def _find_run_indices(paragraph: Paragraph, char_start: int, char_end: int) -> list[int]:
    """Znajduje indeksy runów zawierających placeholder."""
    indices = []
    pos = 0
    for i, run in enumerate(paragraph.runs):
        run_start = pos
        run_end = pos + len(run.text)
        if run_end > char_start and run_start < char_end:
            indices.append(i)
        pos = run_end
    return indices


def get_template_summary(placeholders: list[PlaceholderLocation]) -> dict:
    """Generuje podsumowanie wymaganych pól z szablonu."""
    categories = {}
    for p in placeholders:
        cat = p.category
        if cat not in categories:
            categories[cat] = []
        categories[cat].append({
            'field_id': p.field_id,
            'description': p.description,
            'placeholder': p.placeholder_text[:20],
        })
    return {
        'total_fields': len(placeholders),
        'categories': categories,
    }
